"""
Извлечение текста из .docx для вставки в общий документ без OCR.
"""

from __future__ import annotations

import io

from docx import Document


def extract_text_from_docx(data: bytes) -> str:
    """Параграфы и простые переносы строк; таблицы — построчно через |."""
    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()
