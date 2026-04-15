"""
Постобработка DOCX после Pandoc (границы таблиц).
"""

from __future__ import annotations

import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table


def _set_table_borders(table: Table) -> None:
    """Все границы таблицы и ячеек — чёрные, одинарная линия ~1 pt."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")  # eighths of a point: 8 ≈ 1 pt
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def _iter_all_tables_from_table(table: Table):
    yield table
    for row in table.rows:
        for cell in row.cells:
            for nested in cell.tables:
                yield from _iter_all_tables_from_table(nested)


def _all_tables(doc: Document):
    for t in doc.tables:
        yield from _iter_all_tables_from_table(t)


def apply_black_table_borders(docx_bytes: bytes) -> bytes:
    doc = Document(io.BytesIO(docx_bytes))
    seen: set[int] = set()
    for table in _all_tables(doc):
        tid = id(table._tbl)
        if tid in seen:
            continue
        seen.add(tid)
        _set_table_borders(table)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
